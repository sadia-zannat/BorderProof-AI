"""Generate editable 16:9 competition defense slides using actual synthetic benchmark output."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from PIL import Image
import json

ROOT=Path(__file__).resolve().parents[1]
bench=json.loads((ROOT/'benchmarks/results/holdout_results.json').read_text())
pol=bench['policies'];assert bench['cases']==60
OUT=ROOT/'submissions/BorderProof_AI_FinVeritas_10min_Defense.pptx'
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
BG=RGBColor(8,23,43);PANEL=RGBColor(16,42,66);BORDER=RGBColor(47,82,108)
WHITE=RGBColor(244,250,255);MUTED=RGBColor(160,181,203);TEAL=RGBColor(70,223,192)
BLUE=RGBColor(111,179,250);RED=RGBColor(255,143,136);YELLOW=RGBColor(254,197,120)

def rect(sl,x,y,w,h,fill=PANEL,line=None,radius=False):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=fill
 sh.line.fill.background() if line is None else None
 if line:sh.line.color.rgb=line
 if radius:
  try:sh.adjustments[0]=0.09
  except Exception:pass
 return sh

def tx(sl,x,y,w,h,txt,size=19,color=WHITE,bold=False,align=None):
 shape=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame;tf.word_wrap=True
 tf.margin_left=tf.margin_right=Inches(.02);tf.margin_top=tf.margin_bottom=Inches(.02)
 for n,line in enumerate(str(txt).split('\n')):
  p=tf.paragraphs[0] if n==0 else tf.add_paragraph();p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color;p.space_after=Pt(8)
  if align:p.alignment=align
 return shape

def base(section,title,sub='',n=0):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=BG
 rect(s,0,0,13.333,.06,TEAL)
 tx(s,.62,.25,7,.3,section.upper(),11,TEAL,True)
 tx(s,.62,.68,12.2,.69,title,31,WHITE,True)
 if sub:tx(s,.63,1.44,12.1,.52,sub,15,MUTED)
 rect(s,.62,7.14,12.08,.013,BORDER)
 tx(s,.65,7.18,8,.2,'FINVERITAS  •  BORDERPROOF AI  •  INTERNAL SYNTHETIC SANDBOX',9,MUTED)
 tx(s,11.96,7.18,.65,.2,f'{n:02d}/12',9,MUTED,align=PP_ALIGN.RIGHT)
 return s

def card(s,x,y,w,h,title,body,accent=TEAL):
 rect(s,x,y,w,h,PANEL,BORDER,True)
 rect(s,x+.19,y+.24,.055,.43,accent)
 tx(s,x+.38,y+.24,w-.65,.43,title,18,WHITE,True)
 tx(s,x+.38,y+.89,w-.7,h-1.03,body,15,MUTED)

def note(s,content):
 try:
  frame=s.notes_slide.notes_text_frame
  if frame:frame.text=content
 except Exception:pass

s=base('Shenzhen FinTechathon 2026','BorderProof AI', 'Evidence-consistent student tuition assistance | International AI Topic C',1)
tx(s,.68,2.0,10,1.24,'Verify the evidence.\nProtect the authorization.',36,WHITE,True)
rect(s,.73,4.57,7.4,.76,PANEL,BORDER,True);tx(s,.98,4.76,7.0,.34,'Team FinVeritas  •  English defense  •  10 minutes',18,TEAL,True)
tx(s,.72,5.72,10,.55,'DEMONSTRATION SOFTWARE — NOT A BANK OR LIVE PAYMENT PRODUCT',17,YELLOW,True)
note(s,'Introduce team FinVeritas and explicitly state that all transfers and beneficiaries shown are synthetic; real settlement is not supported.')

s=base('01 / Problem','Tuition payments fail at evidence boundaries','A student can receive a plausible invoice that contains changed beneficiary details.',2)
card(s,.7,2.16,3.8,3.22,'Evidence can conflict','Invoice, recipient and university banking instructions may contradict one another.',BLUE)
card(s,4.76,2.16,3.8,3.22,'Approval can become stale','A quote or beneficiary can change after the student approves a transfer.',YELLOW)
card(s,8.82,2.16,3.8,3.22,'Agent output can mislead','An LLM explanation or forged tool response must not count as payment proof.',RED)
tx(s,.73,5.87,11.8,.62,'Core design principle: untrusted content must never become execution authority.',21,TEAL,True)
note(s,'Real payment providers already exist; this project addresses verifying changing evidence and ensuring authorization remains bound to the intended payment.')

s=base('02 / Competition fit','Why Topic C and why an AI participant?','International students need understandable, auditable cross-border tuition workflows.',3)
card(s,.70,2.05,4,3.08,'TASK COMPLETION · 40%','Working case → invoice → synthetic beneficiary match → quote → approval → simulated transfer → receipt.',TEAL)
card(s,4.85,2.05,3.78,3.08,'SECURITY · 30%','Role/tenant guards, consent binding, fresh evidence checks, negative tests and audit chain.',BLUE)
card(s,8.79,2.05,3.84,3.08,'INNOVATION · 30%','Student-specific BPCE dependency checks and a reproducible synthetic fault benchmark.',YELLOW)
tx(s,.75,5.56,11.5,.77,'Scope: Bangladesh → UK / USA / Canada *demo corridors*; legal eligibility unverified.',18,MUTED)
note(s,'Scoring percentages and topic are taken from the provided international rules. Do not describe the demo corridor list as authorized real-bank routes.')

s=base('03 / Workflow','One auditable tuition-payment journey','The backend enforces control points; the UI cannot bypass them.',4)
stages=[('1','Intent'),('2','Invoice'),('3','Evidence'),('4','Quote'),('5','Approval'),('6','Sandbox'),('7','Receipt')]
for i,(num,label) in enumerate(stages):
 x=.63+i*1.79;rect(s,x,2.5,1.63,1.32,PANEL,BORDER,True);tx(s,x+.14,2.69,1.3,.30,num,21,TEAL,True);tx(s,x+.12,3.12,1.43,.35,label,16,WHITE,True)
 if i<6:tx(s,x+1.61,2.99,.28,.35,'→',19,TEAL,True)
rect(s,.8,4.42,11.76,1.05,PANEL,BORDER,True)
tx(s,1.11,4.67,11.19,.59,'Fail closed on missing evidence, tampering, recipient conflict or expired authorization.',19,WHITE,True)
note(s,'Walk through a single example from the prepared synthetic invoice. Emphasize that approval is case-and-quote specific, not a global permission.')

s=base('04 / Architecture','Real app state, independently enforced security','Frontend is a Next.js source project plus verified FastAPI standalone preview.',5)
for x,head,body in [(.72,'USER UI','Next.js source\nStandalone browser preview\nEN · 简体中文 · বাংলা'),(4.78,'API / BPCE','FastAPI validation\nPolicy and ownership checks\nQuote and consent binding'),(8.83,'DATA / AUDIT','SQLite local demo\nDouble-entry sim ledger\nHash-linked local audit')]:card(s,x,2.04,3.77,3.28,head,body)
tx(s,.8,5.77,11.9,.55,'No financial-provider credentials. No bank settlement. No external identity verification.',19,YELLOW,True)
note(s,'The standalone preview has been exercised with real API requests. The Next.js production build has not been verified because npm registry access was unavailable.')

s=base('05 / BPCE','The BorderProof Consistency Engine','An execution-time recheck, not an assertion from model text.',6)
card(s,.71,2.03,5.76,3.32,'BOUND DEPENDENCIES','Document SHA-256 and parser values\nCase amount / currency\nBeneficiary from synthetic registry\nProvider quote and expiry\nAuthorization scope / evidence-set version',TEAL)
card(s,6.78,2.03,5.79,3.32,'NEGATIVE CONTROLS','Changed amount or recipient → block\nUploaded new evidence → invalidate\nStale quote → reject\nForged simulator response → reject\nUntrusted instructions → flag / stop',RED)
tx(s,.81,5.76,11.8,.64,'Security claims are limited to the tested internal simulator and local threat model.',18,MUTED)
note(s,'BPCE is not cryptographic payment mandate verification. Existing protocols such as AP2 already cover related mandate/authorization ideas.')

s=base('06 / Demonstration','Show one safe completion and one deliberate block','Demonstrate observable backend state and explain failed outcomes.',7)
image_path=ROOT/'docs/V3_BPCE_STUDENT_SCREENSHOT.png'
if image_path.exists():
 im=Image.open(image_path);iw,ih=im.size;wid=6.1;hei=min(3.75,wid*ih/iw)
 s.shapes.add_picture(str(image_path),Inches(.73),Inches(2.07),width=Inches(wid),height=Inches(hei))
card(s,7.2,2.05,5.39,1.68,'PASS','Synthetic correct fixture → explicit approval → COMPLETED simulation.',TEAL)
card(s,7.2,3.94,5.39,1.67,'BLOCK','Changed beneficiary or expired quote → HTTP 409 / no debit.',RED)
tx(s,.82,6.20,11,.34,'Evidence comes from generated DB receipts and reproducible test cases.',15,MUTED)
note(s,'Switch between Payment Lab and BPCE inspector. Never say the example university or bank account has been externally authenticated.')

s=base('07 / StuPay-IntegrityBench','300 authored synthetic scenarios','The holdout set is 60 examples; labels and heuristics are project-authored, not blind expert-validated.',8)
cat=[('Legitimate',60),('Incomplete',40),('Recipient conflict',50),('Prompt injection',50),('Approval changes',40),('Failure & recovery',40),('Corridor conflict',20)]
for i,(lab,val) in enumerate(cat):
 x=.8+(i%4)*3.15;y=2.05+(i//4)*1.84
 rect(s,x,y,2.89,1.54,PANEL,BORDER,True);tx(s,x+.16,y+.16,2.55,.54,str(val),31,TEAL,True);tx(s,x+.16,y+.85,2.66,.5,lab,14,WHITE)
tx(s,.82,6.14,11.8,.32,'Split: 240 development + 60 stratified holdout · 8 hand-coded policy variants.',15,MUTED)
note(s,'The generated benchmark is reproducible, not a random sample of real payments. Held-out labels are not independently audited.')

s=base('08 / Measured results','Unsafe scenarios accepted in 60-case holdout','Lower is better · 48 unsafe-labelled cases · policy proxies are NOT LLMs or AP2.',9)
comparisons=[('Heuristic proxy','heuristic_assistant_proxy',BLUE),('AP2-inspired proxy','ap2_inspired_proxy',YELLOW),('BPCE core proxy','bpce_core_proxy',TEAL)]
for i,(label,key,c) in enumerate(comparisons):
 val=pol[key]['unsafe_acceptances'];y=2.22+i*1.19
 tx(s,.87,y,3.09,.54,label,19,WHITE,True);rect(s,4.03,y,7.13,.41,RGBColor(31,56,78),None,True)
 rect(s,4.03,y,max(.08,val/48*7.13),.41,c,None,True)
 tx(s,11.35,y-.03,1.25,.52,f'{val}/48',22,c,True)
tx(s,.86,6.07,11.9,.63,'BPCE also rejected 3/12 legitimate cases (25% false rejection in this synthetic holdout).',17,RED,True)
note(s,'These numbers come from benchmark results JSON in the ZIP. They are not statistical evidence of real-world finance safety.')

s=base('09 / Failure analysis','Where the current system fails','Measuring faults is more useful than claiming perfect accuracy.',10)
card(s,.72,2.06,3.77,3.3,'50% evasion','Held-out prompt-injection family evaded the heuristic BPCE proxy.',RED)
card(s,4.78,2.06,3.77,3.3,'25% false rejects','Three of 12 legitimate synthetic cases were not completed by the BPCE proxy.',YELLOW)
card(s,8.84,2.06,3.77,3.3,'External trust gap','No real beneficiary attestations, provider settlement or regulatory clearance.',BLUE)
tx(s,.76,5.76,11.9,.7,'Most important next experiment: independently labelled tasks with real comparable agent baselines.',18,TEAL,True)
note(s,'Explicitly acknowledge the inability to generalize. This project is not fraud detection software validated on human cases.')

s=base('10 / Evaluation evidence','Implementation checks and limitations','Evidence categories must not be conflated.',11)
card(s,.75,2.02,5.8,3.24,'REPRODUCIBLE','Pytest API/security suite\nChromium browser interactions\nLocal DB upgrade path\nBenchmark corpus and CSV traces\nDemo export with SHA-256 digest',TEAL)
card(s,6.74,2.02,5.83,3.24,'STILL UNVERIFIED','Next.js production build\nDocker image deployment\nProvider sandbox integrations\nIndependent benchmark labels\nProduction auth / step-up / KYC',YELLOW)
note(s,'Do not give a test count unless the release report is final; show the exact count from current verification. If asked whether the project is deployable, explain the limitations.')

s=base('11 / Closing','Evidence before execution. Accountability after.','A measurable student-finance simulator with a defined path to real validation.',12)
tx(s,.85,2.08,11.65,1.05,'What we delivered',28,TEAL,True)
tx(s,.88,3.05,11.5,1.0,'A working synthetic workflow • BPCE checks • replayable test cases • decision receipts',22,WHITE)
rect(s,.86,4.75,11.5,.9,PANEL,BORDER,True);tx(s,1.14,4.98,11.1,.4,'Not an operational payment provider. Not a verified bank integration.',20,YELLOW,True)
tx(s,.89,6.28,11.34,.39,'FinVeritas  |  BorderProof AI  |  Source and reproducible artifacts in submission ZIP',15,MUTED)
note(s,'Invite questions about BPCE failure modes, baselines, authorization binding and the path to independent data. End in under 10 minutes.')

OUT.parent.mkdir(parents=True,exist_ok=True);prs.save(str(OUT))
reopened=Presentation(str(OUT));assert len(reopened.slides)==12
print('PPTX_CREATED',OUT,'SLIDES',len(reopened.slides),'BYTES',OUT.stat().st_size)
